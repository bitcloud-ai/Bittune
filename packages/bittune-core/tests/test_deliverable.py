"""Phase 7 stage 4: a workbook from sweep facts only, with an empty P/D sheet."""

import io
import json
import sys
import unittest
import zipfile
from pathlib import Path

TESTS_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTS_ROOT.parent / "src"))
sys.path.insert(0, str(TESTS_ROOT))

from openpyxl import load_workbook  # noqa: E402

from bittune_core.contracts import WORKBOOK_BEST_COLUMNS, WORKBOOK_DETAIL_COLUMNS, WORKBOOK_PD_COLUMNS  # noqa: E402
from bittune_core.deliverable import export_deliverable  # noqa: E402
from bittune_core.store import utc_now  # noqa: E402
from test_suite_sweep import MODEL_REF, RUNTIME_SPEC, SweepTests  # noqa: E402


class DeliverableTests(SweepTests):
    def _suite(self) -> str:
        suite_id = "suite-00000000-0000-4000-8000-0000000000d1"
        self.core.store.put_suite_result({
            "schema_version": 1,
            "suite_result_id": suite_id,
            "profile_id": "bittune-default-v1",
            "profile_hash": "sha256:" + "e" * 64,
            "service_instance_id": self.service_id,
            "model_ref": MODEL_REF,
            "runtime_spec": RUNTIME_SPEC,
            "rows": [
                {
                    "scenario_id": "short-chat",
                    "projection": {"scenario": "chat", "input_bucket": "0-2048"},
                    "concurrency": 1,
                    "state": "healthy",
                    "reason_code": "ok",
                    "reason_detail": "fixture",
                    "goodput_req_per_s": 3,
                    "metrics": {"requests_per_second": 3, "output_tokens_per_second": 30, "p50_ttft_ms": 1, "p99_ttft_ms": 2, "p99_e2e_latency_ms": 3, "success_rate": 1},
                },
                {
                    "scenario_id": "short-chat",
                    "projection": {"scenario": "chat", "input_bucket": "0-2048"},
                    "concurrency": 4,
                    "state": "invalid_measurement",
                    "reason_code": "process_failed",
                    "reason_detail": "fixture",
                },
            ],
            "created_at": utc_now(),
            "provenance": "derived",
        }, "run-00000000-0000-4000-8000-000000000000")
        return suite_id

    def _book(self) -> tuple:
        result = self.core.request(export_deliverable, {"intent": "result_workbook", "service_instance_id": self.service_id})
        artifact_id = result["data"]["artifact"]["artifact_id"]
        raw = self.core.store.read_artifact_text(artifact_id)["text"].encode("utf-8", errors="surrogateescape")
        # xlsx is binary. Read the stored file rather than the text view.
        row = self.core.store.db.execute("SELECT relative_path FROM artifacts WHERE artifact_id = ?", (artifact_id,)).fetchone()
        data = (self.core.store.home / row["relative_path"]).read_bytes()
        book = load_workbook(io.BytesIO(data))
        return book, data

    def test_sweep_only_keeps_three_sheets_and_blank_pd(self) -> None:
        self._suite()
        book, _ = self._book()
        self.assertEqual(book.sheetnames, ["全量明细", "最佳吞吐点", "P-D"])
        detail = book["全量明细"]
        self.assertEqual([cell.value for cell in detail[1]], list(WORKBOOK_DETAIL_COLUMNS))
        self.assertEqual(detail.max_row, 3)
        self.assertEqual(detail.cell(3, 4).value, "invalid_measurement")
        self.assertIsNone(detail.cell(3, 6).value)
        best = book["最佳吞吐点"]
        self.assertEqual([cell.value for cell in best[1]], list(WORKBOOK_BEST_COLUMNS))
        self.assertEqual(best.cell(2, 3).value, 3)
        pd_sheet = book["P-D"]
        self.assertEqual([cell.value for cell in pd_sheet[1]], list(WORKBOOK_PD_COLUMNS))
        self.assertEqual(pd_sheet.max_row, 1)

    def test_evidence_pack_hashes_redacted_bytes_and_omits_expected_text(self) -> None:
        self.core.store.put_correctness_result({
            "schema_version": 1,
            "correctness_result_id": "correct-00000000-0000-4000-8000-0000000000d2",
            "profile_id": "bittune-default-v1",
            "service_instance_id": self.service_id,
            "state": "healthy",
            "sample_count": 1,
            "accuracy": 1,
            "samples": [{"sample_id": "s1", "request_id": "chat-1", "passed": True}],
            "api_key": "sk-live-secret",
            "created_at": utc_now(),
            "provenance": "derived",
        }, "run-00000000-0000-4000-8000-000000000000")
        result = self.core.request(export_deliverable, {"intent": "evidence_pack", "service_instance_id": self.service_id})
        artifact_id = result["data"]["artifact"]["artifact_id"]
        row = self.core.store.db.execute("SELECT relative_path, sha256 FROM artifacts WHERE artifact_id = ?", (artifact_id,)).fetchone()
        data = (self.core.store.home / row["relative_path"]).read_bytes()
        from bittune_core.contracts import sha256_bytes

        self.assertEqual(sha256_bytes(data), row["sha256"])
        self.assertEqual(result["data"]["sha256"], row["sha256"])
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            self.assertIn("facts/correctness_result.json", names)
            body = archive.read("facts/correctness_result.json").decode()
            suite_body = archive.read("facts/suite_result.json").decode()
        self.assertNotIn("expected", body)
        self.assertNotIn("sk-live-secret", body)
        self.assertIn("[REDACTED]", body)
        self.assertIn("未测", suite_body)

    def test_profile_id_does_not_fall_back_to_another_profile(self) -> None:
        self._suite()
        self.core.store.put_pd_estimate({
            "schema_version": 1,
            "pd_estimate_id": "pd-00000000-0000-4000-8000-0000000000a1",
            "suite_result_id": "suite-00000000-0000-4000-8000-0000000000d1",
            "profile_id": "profile-a",
            "service_instance_id": self.service_id,
            "state": "healthy",
            "reason_code": "ok",
            "reason_detail": "fixture",
            "p_cap_req_per_s": 4,
            "d_cap_tok_per_s": 8,
            "ratio_p_to_d": 0.5,
            "created_at": "2026-09-01T00:00:00Z",
            "provenance": "derived",
        }, "run-00000000-0000-4000-8000-000000000000")
        self.core.store.put_pd_estimate({
            "schema_version": 1,
            "pd_estimate_id": "pd-00000000-0000-4000-8000-0000000000b1",
            "suite_result_id": "suite-00000000-0000-4000-8000-0000000000d1",
            "profile_id": "profile-b",
            "service_instance_id": self.service_id,
            "state": "healthy",
            "reason_code": "ok",
            "reason_detail": "profile-b-only",
            "p_cap_req_per_s": 9,
            "d_cap_tok_per_s": 8,
            "ratio_p_to_d": 0.5,
            "created_at": "2026-09-02T00:00:00Z",
            "provenance": "derived",
        }, "run-00000000-0000-4000-8000-000000000000")
        def workbook(profile_id: str | None = None):
            params = {"intent": "result_workbook", "service_instance_id": self.service_id}
            if profile_id:
                params["profile_id"] = profile_id
            result = self.core.request(export_deliverable, params)
            artifact_id = result["data"]["artifact"]["artifact_id"]
            row = self.core.store.db.execute("SELECT relative_path FROM artifacts WHERE artifact_id = ?", (artifact_id,)).fetchone()
            return load_workbook(io.BytesIO((self.core.store.home / row["relative_path"]).read_bytes()))

        self.assertEqual(workbook()["P-D"].cell(2, 1).value, 9)
        self.assertEqual(workbook("profile-a")["P-D"].cell(2, 1).value, 4)
        self.assertEqual(workbook("missing-profile")["P-D"].max_row, 1)
        packed = self.core.request(export_deliverable, {"intent": "evidence_pack", "service_instance_id": self.service_id, "profile_id": "missing-profile"})
        packed_id = packed["data"]["artifact"]["artifact_id"]
        packed_row = self.core.store.db.execute("SELECT relative_path FROM artifacts WHERE artifact_id = ?", (packed_id,)).fetchone()
        with zipfile.ZipFile(io.BytesIO((self.core.store.home / packed_row["relative_path"]).read_bytes())) as archive:
            body = archive.read("facts/pd_estimate.json").decode()
        self.assertIn("未测", body)
        self.assertNotIn("profile-b-only", body)


if __name__ == "__main__":
    unittest.main()
