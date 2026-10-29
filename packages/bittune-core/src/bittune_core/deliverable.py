"""Workbook and evidence zip. Both read the store. Neither measures again."""

from __future__ import annotations

import io
import json
import zipfile
from typing import Any

from openpyxl import Workbook

from bittune_core.contracts import (
    WORKBOOK_BEST_COLUMNS,
    WORKBOOK_DETAIL_COLUMNS,
    WORKBOOK_DETAIL_COLUMNS_V2,
    WORKBOOK_PD_COLUMNS,
    ContractError,
    sha256_bytes,
    sha256_canonical,
)
from bittune_core.redaction import redact
from bittune_core.measurement import CONTRACT
from bittune_core.runs import ManagedRun
from bittune_core.store import CoreStore

_DETAIL_METRICS = (
    "requests_per_second",
    "output_tokens_per_second",
    "p50_ttft_ms",
    "p99_ttft_ms",
    "p99_e2e_latency_ms",
    "success_rate",
    "goodput_req_per_s",
)


def export_deliverable(params: dict[str, Any], context: dict[str, Any]) -> dict[str, Any]:
    store: CoreStore = context["store"]
    intent = params.get("intent")
    if intent not in ("result_workbook", "evidence_pack"):
        raise ContractError("invalid_input", "export_deliverable 只接受 result_workbook 或 evidence_pack。", {"intent": intent})
    service_id = params.get("service_instance_id")
    if not isinstance(service_id, str) or not service_id:
        raise ContractError("invalid_input", "export_deliverable 需要 service_instance_id。")
    with ManagedRun(store, "export_deliverable", params, context["session_id"], context) as run:
        profile_id = params.get("profile_id")
        suite_id, benchmark_ids = params.get("suite_result_id"), params.get("benchmark_result_ids")
        if suite_id is not None and benchmark_ids is not None:
            raise ContractError("invalid_input", "suite_result_id与benchmark_result_ids只能选择一个。")
        suite = _latest(store, "suite_result", service_id, profile_id)
        if suite_id is not None:
            suite = store.get_fact("suite_result", suite_id)["payload"]
            if suite.get("service_instance_id") != service_id or (profile_id and suite.get("profile_id") != profile_id):
                raise ContractError("invalid_input", "导出选择与服务/档案不符。")
        if benchmark_ids is not None:
            if not isinstance(benchmark_ids, list) or not benchmark_ids or any(not isinstance(v, str) for v in benchmark_ids):
                raise ContractError("invalid_input", "benchmark_result_ids需要非空ID数组。")
            rows=[]
            for index, bid in enumerate(dict.fromkeys(benchmark_ids)):
                value=store.get_fact("benchmark_result", bid)["payload"]
                if value.get("service_instance_id") != service_id:
                    raise ContractError("invalid_input", "导出实测不属于所选服务。")
                measurement=value.get("measurement") or {}
                rows.append({"scenario_id":"single-point","concurrency":value.get("workload",{}).get("concurrency"),
                             "planned_input_tokens":value.get("workload",{}).get("input_tokens"),
                         "planned_output_tokens":value.get("workload",{}).get("output_tokens"),
                         "planned_request_count":value.get("workload",{}).get("request_count"),
                         "repeat":index+1,"state":"measured" if measurement else "healthy", "reason_code":"explicit_selection",
                             "measurement_status":measurement.get("status"),"slo_status":measurement.get("slo",{}).get("status"),
                             "benchmark_result_id":bid,"measurement":measurement,"metrics":value.get("metrics",{}),
                             "goodput_req_per_s":value.get("metrics",{}).get("goodput_req_per_s")})
            versions={bool(row["measurement"]) for row in rows}
            if len(versions) != 1:
                raise ContractError("invalid_input", "不同测量契约请分别导出，不能混用列语义。")
            suite={"schema_version":2 if True in versions else 1,"view_kind":"explicit_benchmark_selection","rows":rows,"profile_id":None}

        if intent == "result_workbook":
            content = _workbook(suite, _latest(store, "pd_estimate", service_id, profile_id))
            artifact = run.artifact("result-workbook", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            return run.success("交接表已写出。缺的页留空，不重新测量。", {"artifact": artifact})
        content = _pack(store, service_id, suite, profile_id)
        artifact = run.artifact("evidence-pack", content, "application/zip")
        return run.success("证据包已写出。", {"artifact": artifact, "sha256": sha256_bytes(content)})


def _latest(store: CoreStore, kind: str, service_id: str, profile_id: Any) -> dict[str, Any] | None:
    filters: dict[str, Any] = {"service_instance_id": service_id}
    if isinstance(profile_id, str) and profile_id:
        filters["profile_id"] = profile_id
    rows = store.list_facts(kind, limit=1, **filters)
    return rows[0]["payload"] if rows else None


def _workbook(suite: dict[str, Any] | None, estimate: dict[str, Any] | None) -> bytes:
    book = Workbook()
    detail = book.active
    detail.title = "全量明细"
    modern = bool(suite and suite.get("schema_version") == 2)
    detail.append(list(WORKBOOK_DETAIL_COLUMNS_V2 if modern else WORKBOOK_DETAIL_COLUMNS))
    best = book.create_sheet("最佳吞吐点")
    best.append(list(WORKBOOK_BEST_COLUMNS))
    pd_sheet = book.create_sheet("P-D")
    pd_sheet.append(list(WORKBOOK_PD_COLUMNS))
    if suite:
        by_scenario: dict[str, list[dict[str, Any]]] = {}
        for row in suite.get("rows") or []:
            detail.append(_detail_cells_v2(row) if modern else _detail_cells(row))
            by_scenario.setdefault(row.get("scenario_id"), []).append(row)
        for scenario_id, rows in by_scenario.items():
            if modern:
                eligible = [r for r in rows if r.get("state") == "measured" and r.get("slo_status") == "passed"]
                identity = {r.get("measurement", {}).get("dataset_input_sha256") for r in eligible}
                policies = {json.dumps(r.get("measurement", {}).get("acceptance"), sort_keys=True) for r in eligible}
                controlled = eligible and len(policies) == 1 and None not in identity and len(identity) == 1 and all(r.get("measurement", {}).get("cache_policy", {}).get("mode") == "disabled" for r in eligible)
                measured = [r for r in eligible if isinstance(r.get("goodput_req_per_s"), (int, float))]
                chosen = max(measured, key=lambda r: r["goodput_req_per_s"]) if controlled and measured else None
                best.append([scenario_id, chosen["concurrency"] if chosen else None,
                             chosen["goodput_req_per_s"] if chosen else None,
                             chosen["metrics"].get("output_tokens_per_second") if chosen else None,
                             "slo_passed_comparable" if chosen else "selection_unavailable"])
                continue
            healthy = [row for row in rows if row.get("state") == "healthy" and isinstance(row.get("goodput_req_per_s"), (int, float))]
            chosen = max(healthy, key=lambda row: row["goodput_req_per_s"]) if healthy else rows[0]
            passed = chosen.get("state") == "healthy" and bool(healthy)
            best.append([
                scenario_id,
                chosen.get("concurrency") if passed else None,
                chosen.get("goodput_req_per_s") if passed else None,
                (chosen.get("metrics") or {}).get("output_tokens_per_second") if passed else None,
                chosen.get("state"),
            ])
    if estimate and estimate.get("state") == "healthy":
        pd_sheet.append([estimate.get(name) for name in WORKBOOK_PD_COLUMNS])
    if modern:
        source = book.create_sheet("契约与来源")
        source.append(["contract_version", CONTRACT["version"]])
        source.append(["percentile_algorithm", CONTRACT["percentile_algorithm"]])
        source.append(["token_scope", CONTRACT["token_scope"]])
        source.append(["profile_hash", suite.get("profile_hash")])
        source.append(["metric", "unit"])
        for name, unit in CONTRACT["metrics"].items(): source.append([name, unit])
        for sheet in book:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns: sheet.column_dimensions[column[0].column_letter].width = 22
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, float): cell.number_format = "0.000"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _detail_cells_v2(row):
    measurement = row.get("measurement") or {}
    metrics = measurement.get("metrics") or row.get("metrics") or {}
    extra = {"dataset_input_sha256": measurement.get("dataset_input_sha256"),
             "cache_mode": (measurement.get("cache_policy") or {}).get("mode"),
             "metric_sources": json.dumps(measurement.get("metric_sources", {}), ensure_ascii=False),
             "unavailable_reasons": json.dumps(measurement.get("unavailable_reasons", {"row": row.get("reason_detail")}), ensure_ascii=False)}
    return [row.get(key) if key in row else extra.get(key) if key in extra else metrics.get(key) for key in WORKBOOK_DETAIL_COLUMNS_V2]


def _detail_cells(row: dict[str, Any]) -> list[Any]:
    metrics = row.get("metrics") or {}
    cells: list[Any] = []
    for name in WORKBOOK_DETAIL_COLUMNS:
        if name in ("scenario_id", "concurrency", "state", "reason_code"):
            cells.append(row.get(name))
        elif name == "projection":
            cells.append(json.dumps(row.get("projection"), ensure_ascii=False, sort_keys=True) if row.get("projection") else None)
        elif row.get("state") == "healthy" and name in _DETAIL_METRICS:
            cells.append(row.get(name) if name == "goodput_req_per_s" else metrics.get(name))
        else:
            cells.append(None)
    return cells


def _pack(store: CoreStore, service_id: str, suite: dict[str, Any] | None, profile_id: Any) -> bytes:
    facts = {
        "suite_result": None if suite and suite.get("view_kind") else suite,
        "correctness_result": _latest(store, "correctness_result", service_id, profile_id),
        "cache_pair_result": _latest(store, "cache_pair_result", service_id, profile_id),
        "pd_estimate": _latest(store, "pd_estimate", service_id, profile_id),
    }
    service = store.get_fact("service_instance", service_id)["payload"]
    facts["service_instance"] = service
    if service.get("plan_id"):
        plan = store.get_fact("deployment_plan", service["plan_id"])["payload"]
        facts["deployment_plan"] = plan
        if plan.get("hardware_snapshot_id"):
            facts["hardware_snapshot"] = store.get_fact("hardware_snapshot", plan["hardware_snapshot_id"])["payload"]
    run_ids = _run_ids(store, facts)
    benchmarks = []
    for row in (suite or {}).get("rows", []):
        bid = row.get("benchmark_result_id")
        if isinstance(bid, str):
            fact = store.get_fact("benchmark_result", bid)
            benchmarks.append(fact["payload"])
            if fact.get("source_run_id") not in run_ids:
                run_ids.append(fact["source_run_id"])
    listing, manifest, fact_sources, unavailable_runs = [], [], [], []
    blob = io.BytesIO()
    with zipfile.ZipFile(blob, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        def write(path, body):
            archive.writestr(path, body)
            manifest.append({"path": path, "sha256": sha256_bytes(body), "size_bytes": len(body)})
        for kind, payload in facts.items():
            body = json.dumps(redact(payload) if payload else {"state": "未测", "reason": "没有这类 fact。"}, ensure_ascii=False, sort_keys=True).encode()
            path=f"facts/{kind}.json"
            write(path, body)
            if payload and kind in _FACT_IDS:
                fact=store.get_fact(kind,payload[_FACT_IDS[kind]])
                fact_sources.append(_fact_source(path,kind,fact,body))
        if suite and suite.get("view_kind"):
            write("selection.json", json.dumps(suite, ensure_ascii=False, sort_keys=True).encode())
        for result in benchmarks:
            path=f"facts/benchmarks/{result['benchmark_result_id']}.json"
            body=json.dumps(redact(result), ensure_ascii=False, sort_keys=True).encode()
            write(path, body)
            fact=store.get_fact('benchmark_result',result['benchmark_result_id'])
            fact_sources.append(_fact_source(path,"benchmark_result",fact,body))
        for rid in sorted(set(run_ids)):
            try:
                record=store.get_run(rid)
            except ContractError as error:
                if error.code != 'not_found': raise
                unavailable_runs.append({'run_id':rid,'reason':'source_run_not_found'})
            else:
                write(f"runs/{rid}.json",json.dumps(record,ensure_ascii=False,sort_keys=True).encode())
        artifact_ids = {aid for result in benchmarks for aid in result.get("artifact_ids", [])}
        for run_id in run_ids:
            for artifact in store.list_run_artifacts(run_id): artifact_ids.add(artifact["artifact_id"])
        for aid in sorted(artifact_ids):
            data = store.read_artifact_bytes(aid)
            row = store.artifact_row(aid)
            path = f"artifacts/{aid}"
            write(path, data)
            listing.append({"artifact_id": aid, "label": row["label"], "sha256": row["sha256"], "path": path})
        write("artifacts.json", json.dumps(redact(listing), ensure_ascii=False, sort_keys=True).encode())
        write("measurement-contract.json", json.dumps(CONTRACT, ensure_ascii=False, sort_keys=True).encode())
        archive.writestr("manifest.json", json.dumps({"schema_version": 2, "files": manifest, "facts": fact_sources, "completeness": "partial" if unavailable_runs else "complete", "unavailable_runs": unavailable_runs}, ensure_ascii=False, sort_keys=True).encode())
    return blob.getvalue()



def _fact_source(path,kind,fact,body):
    exported_hash=sha256_canonical(json.loads(body))
    return {"path":path,"kind":kind,"id":fact['id'],"content_hash":fact['content_hash'],
            "source_run_id":fact['source_run_id'],"export_content_hash":exported_hash,
            "projection":"identical" if exported_hash==fact['content_hash'] else "redacted"}


_FACT_IDS = {
        "suite_result": "suite_result_id",
        "correctness_result": "correctness_result_id",
        "cache_pair_result": "cache_pair_result_id",
        "pd_estimate": "pd_estimate_id",
        "service_instance": "service_instance_id",
        "deployment_plan": "plan_id",
        "hardware_snapshot": "hardware_snapshot_id",
    }

def _run_ids(store: CoreStore, facts: dict[str, Any]) -> list[str]:
    found: list[str] = []
    for kind, payload in facts.items():
        if not payload:
            continue
        fact_id = payload.get(_FACT_IDS[kind])
        if not isinstance(fact_id, str):
            continue
        source = store.get_fact(kind, fact_id).get("source_run_id")
        if isinstance(source, str) and source not in found:
            found.append(source)
    return found
